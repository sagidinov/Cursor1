#!/usr/bin/env python3
"""Generate a Word letterhead for ООО «РЕСПЕКТСТРОЙ»."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Emu, Mm, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parent
LOGO = ROOT / "assets" / "logo.png"
OUT = ROOT / "фирменный-бланк.docx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
GOLD = RGBColor(0xC4, 0xA3, 0x5A)
MUTED = RGBColor(0x5B, 0x66, 0x73)
INK = RGBColor(0x1B, 0x24, 0x30)


def set_run_font(run, name="Times New Roman", size=12, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_borders(cell, **sides):
    """sides: top/left/bottom/right = (sz, color) in eighths of a point."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge in sides:
            sz, color = sides[edge]
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def no_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def set_paragraph_spacing(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


def add_header(section):
    header = section.header
    header.is_linked_to_previous = False
    for p in header.paragraphs:
        p.clear()

    table = header.add_table(rows=1, cols=3, width=Mm(180))
    table.autofit = False
    no_table_borders(table)
    table.columns[0].width = Mm(24)
    table.columns[1].width = Mm(98)
    table.columns[2].width = Mm(58)

    row = table.rows[0]
    for cell in row.cells:
        shade_cell(cell, "0B1F3A")
        set_cell_margins(cell, top=90, bottom=90, left=80, right=80)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    logo_cell = row.cells[0]
    logo_cell.paragraphs[0].clear()
    p = logo_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(LOGO), width=Mm(18), height=Mm(18))

    brand = row.cells[1]
    p0 = brand.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p0, 0, 0)
    r = p0.add_run("СТРОИТЕЛЬНАЯ КОМПАНИЯ")
    set_run_font(r, "Calibri", 8, True, GOLD)

    p1 = brand.add_paragraph()
    set_paragraph_spacing(p1, 2, 0)
    r = p1.add_run("ООО «РЕСПЕКТСТРОЙ»")
    set_run_font(r, "Calibri", 16, True, RGBColor(255, 255, 255))

    p2 = brand.add_paragraph()
    set_paragraph_spacing(p2, 2, 0)
    r = p2.add_run("Строительство жилых и нежилых зданий")
    set_run_font(r, "Calibri", 8, False, RGBColor(0xC8, 0xD0, 0xD8))

    contacts = row.cells[2]
    p = contacts.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, 0, 0)
    r = p.add_run("+7 (910) 321-00-22")
    set_run_font(r, "Calibri", 11, True, RGBColor(255, 255, 255))
    p = contacts.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, 0, 0)
    r = p.add_run("stroyka3131@yandex.ru")
    set_run_font(r, "Calibri", 9, False, RGBColor(0xE8, 0xE4, 0xD8))
    p = contacts.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, 0, 0)
    r = p.add_run("г. Белгород, ул. Урожайная, д. 1")
    set_run_font(r, "Calibri", 8, False, GOLD)

    gold = header.add_paragraph()
    set_paragraph_spacing(gold, 0, 2)
    goldPr = gold._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C4A35A")
    pBdr.append(bottom)
    goldPr.append(pBdr)

    meta = header.add_paragraph()
    set_paragraph_spacing(meta, 4, 0)
    r = meta.add_run("ИНН 3123486273  ·  КПП 312301001  ·  ОГРН 1223100002124  ·  ОКПО 52541432")
    set_run_font(r, "Calibri", 8, False, MUTED)


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        p.clear()

    line = footer.paragraphs[0]
    set_paragraph_spacing(line, 0, 4)
    pPr = line._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "12")
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), "C4A35A")
    pBdr.append(top)
    pPr.append(pBdr)

    table = footer.add_table(rows=1, cols=2, width=Mm(180))
    table.autofit = False
    no_table_borders(table)
    table.columns[0].width = Mm(100)
    table.columns[1].width = Mm(80)

    left = table.rows[0].cells[0]
    p = left.paragraphs[0]
    set_paragraph_spacing(p, 0, 0)
    r = p.add_run("ООО «РЕСПЕКТСТРОЙ»")
    set_run_font(r, "Calibri", 8, True, NAVY)
    p = left.add_paragraph()
    set_paragraph_spacing(p, 0, 0)
    r = p.add_run("308010, Белгородская область, г. Белгород,\nул. Урожайная, д. 1, помещ. 2")
    set_run_font(r, "Calibri", 8, False, MUTED)

    right = table.rows[0].cells[1]
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, 0, 0)
    r = p.add_run("ИНН/КПП 3123486273 / 312301001")
    set_run_font(r, "Calibri", 8, False, MUTED)
    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, 0, 0)
    r = p.add_run("ОГРН 1223100002124")
    set_run_font(r, "Calibri", 8, False, MUTED)
    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, 0, 0)
    r = p.add_run("тел. +7 (910) 321-00-22  ·  stroyka3131@yandex.ru")
    set_run_font(r, "Calibri", 8, False, MUTED)


def add_body(doc):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    no_table_borders(table)
    table.columns[0].width = Mm(95)
    table.columns[1].width = Mm(85)

    left = table.rows[0].cells[0]
    p = left.paragraphs[0]
    set_paragraph_spacing(p, 6, 4)
    r = p.add_run("Исх. № ________________")
    set_run_font(r, "Times New Roman", 12, False, INK)
    p = left.add_paragraph()
    set_paragraph_spacing(p, 0, 4)
    r = p.add_run("от «____» ______________ 20____ г.")
    set_run_font(r, "Times New Roman", 12, False, INK)
    p = left.add_paragraph()
    set_paragraph_spacing(p, 0, 8)
    r = p.add_run("На № ________________ от ____________")
    set_run_font(r, "Times New Roman", 12, False, INK)

    right = table.rows[0].cells[1]
    set_cell_borders(right, left=(8, "C4A35A"))
    set_cell_margins(right, top=60, bottom=40, left=140, right=40)
    p = right.paragraphs[0]
    set_paragraph_spacing(p, 0, 4)
    r = p.add_run("АДРЕСАТ")
    set_run_font(r, "Calibri", 8, True, GOLD)
    for _ in range(3):
        p = right.add_paragraph()
        set_paragraph_spacing(p, 2, 2)
        r = p.add_run("________________________________")
        set_run_font(r, "Times New Roman", 12, False, INK)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, 16, 12)
    r = title.add_run("О ________________________________")
    set_run_font(r, "Times New Roman", 13, True, NAVY)

    intro = doc.add_paragraph()
    set_paragraph_spacing(intro, 0, 8, line=1.5)
    r = intro.add_run("Уважаемый _______________________________!")
    set_run_font(r, "Times New Roman", 13, False, INK)

    for _ in range(8):
        p = doc.add_paragraph()
        set_paragraph_spacing(p, 0, 8, line=1.5)
        r = p.add_run("")
        set_run_font(r, "Times New Roman", 13, False, INK)

    sign = doc.add_table(rows=1, cols=3)
    sign.autofit = False
    no_table_borders(sign)
    sign.columns[0].width = Mm(45)
    sign.columns[1].width = Mm(70)
    sign.columns[2].width = Mm(65)

    c0 = sign.rows[0].cells[0]
    p = c0.paragraphs[0]
    set_paragraph_spacing(p, 18, 0)
    r = p.add_run("Директор")
    set_run_font(r, "Times New Roman", 12, False, INK)

    c1 = sign.rows[0].cells[1]
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, 18, 0)
    r = p.add_run("____________________")
    set_run_font(r, "Times New Roman", 12, False, INK)
    p = c1.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, 0, 0)
    r = p.add_run("подпись")
    set_run_font(r, "Calibri", 8, False, MUTED)

    c2 = sign.rows[0].cells[2]
    p = c2.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(p, 18, 0)
    r = p.add_run("Б. Д. Чачашвили")
    set_run_font(r, "Times New Roman", 12, False, INK)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(18)
    section.right_margin = Mm(16)
    section.top_margin = Mm(12)
    section.bottom_margin = Mm(16)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)

    add_header(section)
    add_footer(section)
    add_body(doc)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = INK

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
